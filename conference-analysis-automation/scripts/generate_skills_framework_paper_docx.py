from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.shape import WD_INLINE_SHAPE
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/liziyu/llm_conference_skill")
OUTPUT_DIR = ROOT / "generated-data" / "reports"
DOCX_PATH = OUTPUT_DIR / "skills_framework_paper_cn.docx"
MD_PATH = OUTPUT_DIR / "skills_framework_paper_cn.md"
FIGURE_DIR = OUTPUT_DIR / "figures"
FIG1_PATH = FIGURE_DIR / "figure1_framework.png"
FIG2_PATH = FIGURE_DIR / "figure2_workflow.png"
FIG3_PATH = FIGURE_DIR / "figure3_coverage_metrics.png"


TITLE = "面向顶级学术会议数据治理与分析的基于 Skill 的自动化框架研究"
SUBTITLE = "七大会 2020—2025 数据治理、统计分析与质量审计系统论文初稿"


ABSTRACT = (
    "围绕顶级学术会议的数据采集、标准化、统计分析与质量审计，本文设计并实现了一套基于 Skill "
    "的自动化框架。该框架以模块化方式组织论文抽取、作者与机构解析、国家识别、会议征文信息获取、"
    "统计分析、结果输出与质量核查等处理流程，目标是在统一的数据组织方式下提高多会议、多年份批量处理"
    "的稳定性、可复用性与可扩展性。本文以密码学与网络安全领域七个代表性会议 2020—2025 年的数据为实验对象，"
    "共覆盖 42 个 conference-year 实例。实验结果表明，当前框架已经完成 42 个 conference-year 的论文与会议信息"
    "主链路处理，其中 CFP 主链路在当前范围内已全部打通；在论文级覆盖审计层，最新快照共获得 7426 篇论文，其中 7183 篇"
    "具备作者行记录，论文级作者信息获取率达到 96.73%，具备完整 affiliation 覆盖的论文为 6023 篇，对应论文完整率为 81.11%。"
    "同时，历史全量输出与当前标准化分析层在论文数、作者机构行数、机构数和未知国家行数等方面保持一致，说明该系统的"
    "标准化层已经具有较好的稳定性。进一步分析表明，近期指标提升主要来自论文级覆盖审计与对齐修复，而不是简单扩大下游统计口径。"
    "本文的工作说明，将复杂科研数据处理流程拆分为职责明确的模块，并在统一中间数据组织下进行编排，能够较好地支持顶会数据治理与分析任务的持续迭代。"
)

KEYWORDS = "顶会数据；科研数据治理；自动化分析；模块化框架；作者机构解析；会议征文抽取"


SECTIONS: list[tuple[str, list[str]]] = [
    (
        "1 引言",
        [
            "近年来，围绕顶级学术会议开展的数据收集与分析工作逐渐增多。无论是研究学科发展趋势、地区与机构参与格局，还是分析作者合作关系与论文主题变化，都需要较为完整且结构统一的基础数据。然而，这类工作在实际开展过程中往往面临数据来源分散、页面结构差异大、作者与机构信息噪声多、统计口径容易漂移等问题。",
            "会议官网、开放学术平台、PDF 页面和历史归档往往使用不同字段结构，同一会议在不同年份的组织方式也可能完全不同。进一步地，作者与机构的解析还会受到脚注、通讯作者说明、单位简称、多语言写法和学院实验室层级差异的影响。如果缺少统一的中间数据与持续审计机制，下游统计结果即使形式完整，也未必能够真实反映顶会参与情况。",
            "针对这些问题，本文将顶会数据自动化处理过程组织为一套基于 Skill 的模块化框架。这里的 Skill 更接近职责清晰、可独立执行、又能够接入统一处理流程的功能模块，而不是泛化的能力口号。围绕这一思路，本文将整个系统拆分为采集、标准化、质量增强、统计分析、结果表达和自动化编排等层次，使其既能支撑主数据链路，也能持续承接科研分析任务。",
            "本文以密码学与网络安全方向七个代表性会议 2020—2025 年数据为实验对象，共覆盖 42 个 conference-year 实例。系统重点回答三个问题：一是模块化组织方式是否能够支撑多会议、多年份处理流程稳定运行；二是哪些关键模块会直接影响论文覆盖率、作者信息获取率、机构国家解析率和论文完整率；三是在主数据链路稳定后，统计分析、角色分析与 CFP 信息提取等扩展模块如何提升整体实用性。",
            "本文的主要贡献包括：提出一套面向顶级学术会议数据治理与分析的基于 Skill 的自动化框架；在七大会 2020—2025 全量范围内构建统一中间数据层，并验证历史全量输出与当前标准化分析层的一致性；结合论文级覆盖审计结果，说明覆盖修复与对齐逻辑对核心指标提升的实际作用，为后续系统论文撰写、版本对比与实验复现提供基础。"
        ],
    ),
    (
        "2 相关工作",
        [
            "与本文相关的工作主要包括学术元数据采集、科研评价与学术分析系统，以及面向科研自动化的模块化工作流研究。现有元数据采集工作通常围绕论文题目、作者、机构、DOI、引用关系等字段展开，常见数据源包括会议官网、数字图书馆和开放学术索引平台。这些平台能够提供较好的基础入口，但在作者隶属、机构标准化与关键会议信息方面仍存在字段不完整、口径不一致和跨年份差异明显的问题。",
            "科研评价与学术分析系统更关注国家、机构、作者、主题和合作关系等下游分析输出，能够支持趋势分析、合作网络构建和多维统计展示。但这类工作通常默认上游数据已经相对干净，较少详细展开原始数据如何治理、不同来源如何补齐以及质量问题如何被持续审计。因此，在面向顶会长期批量分析时，上游数据治理的过程本身仍值得单独研究。",
            "近年来，围绕科研写作和科研自动化的模块化工作流研究逐渐增多。相比单体脚本，将复杂流程拆分为职责明确的模块，有利于替换局部策略、记录中间结果并开展版本对比。本文的工作与这一思路一致，但关注对象更加具体，即顶会数据治理与分析这一兼具数据工程和科研分析特征的场景。",
            "与已有工作相比，本文的特点在于：一方面，将顶会数据治理过程本身作为研究对象，不只关注结果展示，也关注抽取、标准化、审计和修复过程如何协同；另一方面，明确区分直接影响数据质量的核心模块和主要扩展分析维度的功能模块，从而使系统设计、实验组织与论文叙述保持清晰一致。"
        ],
    ),
    (
        "3 系统框架设计",
        [
            "本文框架的设计围绕四个目标展开：保证跨会议与跨年份处理流程的统一性；提升数据治理过程的可审计性；支持模块级替换与增量扩展；兼顾主数据链路与下游分析任务。整体上，系统可以划分为采集层、标准化层、质量增强层、统计分析层、结果表达层和自动化编排层。",
            "采集层负责从官网、开放学术平台、PDF 页面和历史补录结果中获取论文、作者、机构与会议元数据。标准化层负责统一会议命名、作者字段、机构文本与国家字段，并组织结构化中间结果。质量增强层主要处理作者与机构错配、噪声行过滤、未知国家识别、论文级覆盖审计和历史数据回填等问题。统计分析层基于标准化后的事实表生成国家、机构、作者、趋势和合作等分析结果。结果表达层负责报告、表格、图表和基准比较文件的生成，自动化编排层则承接批量执行、结果归档、日志组织和版本管理。",
            "按照功能特点，系统模块进一步分为两组。第一组是质量增强型模块，包括抽取模块、补源模块、去噪模块、机构标准化模块、国家识别与规则治理模块、覆盖修复模块、历史补齐模块和质量审计模块。这些模块直接影响论文覆盖率、作者信息获取率、机构国家解析情况以及论文完整率。第二组是分析扩展型模块，包括国家统计、机构统计、作者统计、趋势分析、合作分析、作者角色分析、CFP 提取、报告生成、可视化和基准比较等模块。这类模块主要扩展系统输出能力，不应被简单写成同类性能增强项。",
            "在数据流程上，系统先完成论文与作者信息获取，再进入标准化、质量增强和统计分析流程，最后将当前结果、底层数据、历史归档和运行日志分别组织归档。这样的文件组织方式便于区分当前可用结果与历史实验快照，也便于开展版本对比和实验复现。"
        ],
    ),
    (
        "4 系统实现",
        [
            "系统当前面向七个会议 2020—2025 年数据开展批量处理，共覆盖 42 个 conference-year 实例。处理对象主要包括论文级元数据、作者级记录、机构级记录、国家级映射结果以及会议征文与关键日期信息。其中国家、机构和作者三类数据是后续统计分析的核心基础，CFP 信息则主要服务于会议元数据补齐与时间维度分析。",
            "在数据来源上，系统并不依赖单一平台，而是采用多来源协同的方式获取论文与作者数据。对于不同会议与年份，系统可根据数据可用性使用官网页面、开放学术索引平台、PDF 页面信息以及历史补录结果等来源。这样做的原因在于，不同来源的稳定字段并不相同，单一来源很难同时兼顾论文列表完整性与作者机构完整性。",
            "标准化与质量增强是整个系统中最关键的实现部分。其核心工作包括：机构文本清洗与标准化；国家识别与未知项治理；作者与隶属噪声过滤；论文级覆盖审计与对齐修复；历史补齐与回填机制。实践中，系统重点处理了误将学院、实验室、国家地区、通讯作者说明、脚注或 affiliation 标签识别为作者的情况，并通过规则治理减少这类误识别对作者计数与机构解析的干扰。",
            "在统计分析方面，系统已经能够输出国家、机构、作者、趋势和合作等多类分析结果，并生成多种报表与对比文件。作者角色分析模块当前已支持第一作者、参与作者和共同作者等基础角色的抽取与统计，但通讯作者链路尚未形成稳定主流程，因此本文不将其视为已成熟完成的能力。",
            "需要特别说明的是，作者角色模块当前并不是覆盖 7426 篇论文的全量主结果模块，而是建立在已经进入规范化作者行流水线的数据之上，用于提供角色分析能力的扩展输出。因此，它更适合作为非全量、功能扩展型模块来描述，而不应与主批次覆盖率指标直接并列。",
            "在 CFP 方面，系统已经形成单独的数据处理链路，用于抽取投稿截止日期、通知日期和会议日期等元数据。当前在 42 个 conference-year 范围内，event 与 edition 记录已全部建立，覆盖状态均已达到 complete_unverified，review queue 也已清零，说明该主链路在本文实验范围内已经基本跑通。",
            "为了支持持续迭代，系统还将自动化编排与结果归档作为单独能力建设。当前输出已按当前结果、底层数据、历史归档和运行日志四类进行整理，并辅以验证脚本对主要生成文件进行一致性检查。这使实验结果不再只是单次运行日志，而是可以被追溯、对比和复用的版本化产物。"
        ],
    ),
    (
        "5 实验设计与评价指标",
        [
            "本文实验主要服务于三个目标：验证全量框架在七大会 2020—2025 范围内的可运行性；观察关键质量增强模块对主指标的影响；展示在主数据链路稳定后，系统如何继续扩展分析任务。",
            "实验范围覆盖 ASIACRYPT、CRYPTO、EUROCRYPT、USENIX Security、NDSS、IEEE S&P 和 ACM CCS 七个会议 2020—2025 年数据，共 42 个 conference-year 实例。当前主实验快照显示，系统共处理论文 7426 篇，扩展作者行 31699 条，完成分类作者行 31388 条，单国成功解析作者行 25241 条，未解析国家作者行 311 条。",
            "需要区分的是，标准化层中的“行级字段完整性”和覆盖审计中的“论文级完整性”并不是同一概念。当前标准化事实表中的作者名和机构名字段未出现空值，但这并不意味着论文级缺失已经消失。论文级仍存在没有作者机构行或仅有部分 affiliation 覆盖的情况，需要通过覆盖审计指标单独观察。",
            "结合系统特点，本文采用论文级覆盖指标、作者与机构解析指标、会议级覆盖指标、稳定性指标以及功能完成度指标进行评估。论文级覆盖指标重点观察论文总数、具备作者行记录的论文数、具备完整 affiliation 覆盖的论文数以及 complete_unverified 论文数；稳定性指标则用于比较历史全量输出与当前标准化分析层的差异，判断主事实表是否保持一致。",
            "在当前主快照中，系统覆盖 42 个 conference-year，获得 7426 篇论文，其中 7183 篇具备作者行记录，论文级作者信息获取率达到 96.73%。具备完整 affiliation 覆盖的论文为 6023 篇，占全部论文的 81.11%；complete_unverified 论文数同样为 6023，占比 81.11%。从国家解析情况看，当前扩展作者行数为 31699，完成分类作者行数为 31388，单国成功解析作者行数为 25241，未解析国家作者行数为 311。",
            "与此同时，论文级覆盖审计也表明，当前仍有 243 篇论文没有作者机构行，另有 1160 篇论文仅达到部分 affiliation 覆盖。因此，“标准化行中没有空作者名或空机构名”只能说明已入表记录本身没有空字段，不能替代论文级完整性判断。",
            "为了观察近期指标变化来源，本文对历史全量输出、当前标准化分析层以及近期不同快照的论文级覆盖审计结果进行比较。结果显示，历史全量输出与当前标准化分析层在抽取论文数、标准化论文数、作者机构行数、唯一作者数、唯一机构数和未知国家行数等指标上完全一致，说明当前系统并没有通过改变统计口径来制造提升，而是保持了主事实表的稳定。",
            "进一步地，从论文级覆盖审计快照演化看，较早快照中具备作者行记录的论文数为 6884，具备完整 affiliation 覆盖的论文数为 5490，complete_unverified 论文数为 5379；在最新快照中，这三项指标上升到 7183、6023 和 6023。该结果说明，近期可量化改进主要来自论文级覆盖审计、论文与作者行的对齐修复以及缺失记录识别方式的优化。",
            "除主数据链路外，系统已经能够输出国家、机构、作者、趋势和合作等多类分析结果，并生成多种报表与对比文件。与此同时，第一作者与参与作者等角色分析模块已经具备独立输出能力，CFP 模块也补足了会议关键时间信息。这一部分更适合作为框架扩展能力展示，而不应简单理解为提高抽取性能的同类消融实验。"
        ],
    ),
    (
        "6 讨论",
        [
            "从实验结果看，标准化层长期保持稳定，而论文级覆盖指标在近期明显上升。这说明系统此前的主要瓶颈并不在机构标准化事实表本身，而在于论文级完整性识别逻辑。如果论文与作者行之间的对齐不够充分，或者对缺失与部分缺失的判断不够细，就可能低估完整论文数。当前版本通过补足这部分逻辑，使完整率更接近真实情况。",
            "本文采用的模块化组织方式使不同问题能够在相对独立的层次上被处理。抽取模块关注能否获取数据，标准化模块关注如何统一口径，质量增强模块关注数据是否可信，分析模块关注能做哪些研究问题。这种分层既降低了系统迭代时的混乱程度，也让后续按模块开展版本对比和论文实验设计成为可能。",
            "虽然本文实验范围集中在密码学与网络安全会议，但框架本身并不依赖某一特定会议，只要能够为上游采集建立适配逻辑，就可以将相同的处理中间层和审计逻辑迁移到其他学术方向。对于需要长期维护会议数据集，并希望将结果直接用于报告、统计和论文写作的场景，这种框架具有较好的适用性。",
            "需要说明的是，系统目前仍存在若干限制。第一，通讯作者主链路尚未成熟，不适合写成稳定完成的能力。第二，尽管未知国家行数已经显著压缩，但仍保留少量边界项，需要依赖后续机构治理继续消化。第三，部分会议年份的原始页面结构差异较大，虽然已经能通过补源与回退机制覆盖，但来源一致性仍不完全相同。第四，本文当前实验主要围绕七大会 2020—2025 展开，更长年份跨度的数据仍可继续补齐。",
            "从科研写作角度看，本文的经验说明，将复杂工程工作包装成系统研究时，关键不是罗列功能点，而是清楚地区分主链路增强能力与功能扩展能力，并围绕真正影响主指标的模块组织实验。只有这样，系统设计、实验结果和论文叙述才能形成互相支撑的整体。"
        ],
    ),
    (
        "7 结论",
        [
            "本文围绕顶级学术会议数据采集、治理与分析问题，提出并实现了一套基于 Skill 的自动化框架。该框架通过采集、标准化、质量增强、统计分析、结果表达和自动化编排等多个层次，将原本分散的数据处理脚本组织为可持续迭代的处理流程。",
            "在七大会 2020—2025、共 42 个 conference-year 的实验范围内，系统已经完成论文、作者、机构、国家和会议信息的主链路处理。实验结果表明，当前主快照的论文级作者信息获取率达到 96.73%，论文完整率达到 81.11%；与此同时，历史全量输出与当前标准化分析层结果保持一致，说明标准化层已具备较好的稳定性。进一步比较发现，近期指标提升主要来自覆盖审计和对齐修复，而不是统计口径扩大。",
            "总体来看，本文工作说明，将科研数据治理任务组织为职责明确的模块体系，并在统一中间数据基础上开展自动化编排，是一条可行且具有实践价值的路径。未来工作将继续围绕未知机构治理、作者角色增强、更长时间跨度的历史补齐以及更丰富的自动化分析输出展开。"
        ],
    ),
]


TABLE_1 = [
    ["指标", "数值"],
    ["会议范围", "ASIACRYPT、CRYPTO、EUROCRYPT、USENIX Security、NDSS、IEEE S&P、ACM CCS"],
    ["年份范围", "2020—2025"],
    ["conference-year 数量", "42"],
    ["论文总数", "7426"],
    ["扩展作者行数", "31699"],
    ["标准化作者机构行数", "32527"],
    ["唯一作者数", "12795"],
    ["唯一机构数", "6394"],
]

TABLE_2 = [
    ["层次", "主要模块", "作用"],
    ["采集层", "论文抓取、作者与机构获取、会议信息获取", "从多来源获得论文、作者、机构和会议元数据"],
    ["标准化层", "字段统一、机构清洗、国家字段整理", "形成统一中间数据，降低来源差异对下游影响"],
    ["质量增强层", "补源、去噪、机构治理、国家识别、覆盖修复、历史回填、质量审计", "直接改善覆盖率、完整率与解析稳定性"],
    ["统计分析层", "国家统计、机构统计、作者统计、趋势分析、合作分析、角色分析", "生成面向科研分析的多维结果"],
    ["结果表达层", "报告、表格、图表、基准比较", "将分析结果组织为可阅读、可复用的产物"],
    ["自动化编排层", "批量执行、归档、日志与版本管理", "支撑持续迭代、版本对比和实验复现"],
]

TABLE_3 = [
    ["指标", "历史全量输出", "当前标准化分析层", "差值"],
    ["conference-year 覆盖范围", "42/42", "42/42", "0"],
    ["抽取论文数", "7426", "7426", "0"],
    ["标准化论文数", "6897", "6897", "0"],
    ["标准化作者机构行数", "32527", "32527", "0"],
    ["唯一作者数", "12795", "12795", "0"],
    ["唯一机构数", "6394", "6394", "0"],
    ["未知国家行数", "316", "316", "0"],
]

TABLE_4 = [
    ["快照", "论文总数", "有作者行论文数", "完整 affiliation 覆盖论文数", "完整论文数", "作者信息获取率", "论文完整率"],
    ["早期快照 1", "7426", "6884", "5490", "5379", "92.70%", "72.43%"],
    ["早期快照 2", "7426", "6884", "5490", "5379", "92.70%", "72.43%"],
    ["当前快照", "7426", "7183", "6023", "6023", "96.73%", "81.11%"],
]

TABLE_5 = [
    ["会议", "年份数", "论文数", "完整覆盖论文数", "完整率", "单国解析成功率", "未解析国家行数"],
    ["ASIACRYPT", "6", "654", "643", "98.32%", "69.00%", "14"],
    ["CCS", "6", "1755", "804", "45.81%", "98.18%", "10"],
    ["CRYPTO", "6", "710", "653", "91.97%", "73.63%", "18"],
    ["EUROCRYPT", "6", "580", "426", "73.45%", "98.58%", "19"],
    ["IEEE S&P", "6", "1082", "1001", "92.51%", "96.75%", "89"],
    ["NDSS", "6", "704", "651", "92.47%", "97.35%", "63"],
    ["USENIX Security", "6", "1941", "1845", "95.05%", "97.80%", "98"],
]


def set_cell_text(cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = tbl_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tbl_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "B8C2CC")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_widths(table, widths_cm: list[float]):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def add_table_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def add_table(doc: Document, caption: str, rows: list[list[str]], widths_cm: list[float]):
    add_table_caption(doc, caption)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_table_widths(table, widths_cm)
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or len(text) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(table.cell(r_idx, c_idx), text, bold=(r_idx == 0), align=align)
            if r_idx == 0:
                shade_cell(table.cell(r_idx, c_idx), "F4F6F9")
    doc.add_paragraph()


def apply_run_font(run, name="Calibri", east_asia="宋体", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def ensure_style(document: Document, style_name: str, style_type, base=None):
    styles = document.styles
    if style_name in styles:
        return styles[style_name]
    style = styles.add_style(style_name, style_type)
    if base:
        style.base_style = styles[base]
    return style


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.33
    normal.paragraph_format.first_line_indent = Cm(0.74)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.first_line_indent = Cm(0)

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.first_line_indent = Cm(0)

    h3 = doc.styles["Heading 3"]
    h3.font.name = "Calibri"
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.first_line_indent = Cm(0)

    caption = ensure_style(doc, "CaptionCN", WD_STYLE_TYPE.PARAGRAPH, base="Normal")
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(10.5)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)


def add_title_block(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(TITLE)
    apply_run_font(run, size=20, bold=True, color=(0, 0, 0))

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(14)
    run2 = p2.add_run(SUBTITLE)
    apply_run_font(run2, size=11, color=(89, 89, 89))

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(18)
    run3 = info.add_run("研究范围：七个密码学与网络安全顶会（2020—2025）")
    apply_run_font(run3, size=10.5, color=(89, 89, 89))


def add_abstract(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    label = p.add_run("摘要：")
    apply_run_font(label, east_asia="黑体", bold=True, size=11)
    body = p.add_run(ABSTRACT)
    apply_run_font(body, size=11)

    p2 = doc.add_paragraph()
    p2.paragraph_format.first_line_indent = Cm(0)
    key = p2.add_run("关键词：")
    apply_run_font(key, east_asia="黑体", bold=True, size=11)
    body2 = p2.add_run(KEYWORDS)
    apply_run_font(body2, size=11)


def add_paragraph(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    apply_run_font(run, size=11)


def add_figure_placeholder(doc: Document, caption: str, note: str):
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    set_table_widths(box, [16.0])
    set_table_borders(box)
    cell = box.cell(0, 0)
    shade_cell(cell, "FAFAFA")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(note)
    apply_run_font(run, size=10.5, color=(96, 96, 96))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cap = doc.add_paragraph(style="CaptionCN")
    run2 = cap.add_run(caption)
    apply_run_font(run2, size=10.5, bold=True)
    doc.add_paragraph()


def add_figure(doc: Document, caption: str, image_path: Path, fallback_note: str, width_cm: float = 16.2):
    if not image_path.exists():
        add_figure_placeholder(doc, caption, fallback_note)
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))

    cap = doc.add_paragraph(style="CaptionCN")
    run2 = cap.add_run(caption)
    apply_run_font(run2, size=10.5, bold=True)
    doc.add_paragraph()


def add_references(doc: Document):
    doc.add_heading("参考文献（待补充）", level=1)
    refs = [
        "会议官网、开放学术平台与历史归档中的相关数据来源。",
        "学术元数据采集、机构标准化与科研分析系统方面的相关研究文献。",
        "顶会统计分析与科研自动化工作流相关研究文献。",
    ]
    for idx, item in enumerate(refs, 1):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(f"[{idx}] {item}")
        apply_run_font(run, size=11)


def build_doc():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_title_block(doc)
    add_abstract(doc)

    for title, paragraphs in SECTIONS:
        doc.add_heading(title, level=1)
        for para in paragraphs:
            add_paragraph(doc, para)
        if title == "3 系统框架设计":
            add_figure(
                doc,
                "图 1 基于 Skill 的顶会数据治理与分析自动化框架",
                FIG1_PATH,
                "此处可插入系统总览图：采集层 → 标准化层 → 质量增强层 → 统计分析层 → 结果表达层 → 自动化编排层",
            )
        if title == "4 系统实现":
            add_figure(
                doc,
                "图 2 顶会数据自动化处理流程",
                FIG2_PATH,
                "此处可插入流程图：论文获取 → 作者与机构提取 → 标准化 → 质量审计 → 统计输出 → 归档",
            )
        if title == "5 实验设计与评价指标":
            add_table(doc, "表 1 实验数据集总体情况", TABLE_1, [4.2, 11.8])
            add_table(doc, "表 2 系统模块分层及其主要功能", TABLE_2, [2.2, 5.0, 8.8])
            add_table(doc, "表 3 历史全量输出与当前标准化分析层对比", TABLE_3, [4.0, 3.8, 4.2, 2.0])
            add_table(doc, "表 4 不同版本快照下论文级覆盖审计指标对比", TABLE_4, [2.6, 1.8, 2.5, 3.0, 1.9, 2.4, 1.9])
            add_table(doc, "表 5 七个顶级会议的会议级结果比较", TABLE_5, [3.6, 1.3, 1.6, 2.0, 1.5, 2.1, 1.8])
            add_figure(
                doc,
                "图 3 不同版本快照下论文级覆盖审计指标变化",
                FIG3_PATH,
                "此处可插入柱状图或折线图，对比 papers_with_author_rows、papers_with_full_affiliation_coverage 与 papers_complete_unverified 的变化",
            )

    add_references(doc)

    section = doc.sections[-1]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("基于 Skill 的顶会数据治理与分析自动化框架研究")
    apply_run_font(footer_run, size=9, color=(120, 120, 120))

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "系统论文初稿"
    doc.core_properties.author = "OpenAI Codex"
    doc.save(str(DOCX_PATH))


if __name__ == "__main__":
    build_doc()
    print(DOCX_PATH)
