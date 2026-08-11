from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

fonts = [
    "SimSun", "宋体", "NSimSun", "Songti SC", "STSong", "PingFang SC",
    "Heiti SC", "STHeiti", "Hiragino Sans GB", "Arial Unicode MS",
    "Microsoft YaHei", "FangSong_GB2312",
]
doc = Document()
for name in fonts:
    p = doc.add_paragraph()
    run = p.add_run(f"{name}: 学术会议论文元数据处理测试 ABC 123")
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr, val in (("ascii", "Times New Roman"), ("hAnsi", "Times New Roman"),
                      ("cs", "Times New Roman"), ("eastAsia", name), ("hint", "eastAsia")):
        rfonts.set(qn(f"w:{attr}"), val)
doc.save("/tmp/font-probe.docx")
