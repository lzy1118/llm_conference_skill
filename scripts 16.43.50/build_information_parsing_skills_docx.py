from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/liziyu/llm_conference_skill")
OUTPUT = Path("/Users/liziyu/Documents/upgrade/信息解析Skills汇编.docx")

SOURCES = [
    (
        "conference-paper-extraction",
        "原始论文数据解析",
        "从 papers.json 等原始数据生成论文、作者与机构文本记录。",
    ),
    (
        "conference-grobid-extract-metadata",
        "PDF/TEI 元数据解析",
        "从 PDF 和 GROBID TEI 中解析标题、作者、机构、邮箱、摘要及作者角色证据。",
    ),
    (
        "conference-clean-paper-metadata",
        "论文—作者元数据清洗",
        "审计、修复并补全解析后的论文—作者元数据。",
    ),
    (
        "conference-author-role-extraction",
        "作者角色事实抽取",
        "生成第一作者、参与作者、通讯作者和共同一作等角色事实及复核标记。",
    ),
]

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "243447"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CODE_FILL = "F6F8FA"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=11, bold=None, italic=None, color=INK, latin="Microsoft YaHei", east_asia="Microsoft YaHei"):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Title": (28, DARK_BLUE, 0, 8),
        "Subtitle": (14, MUTED, 0, 10),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
        "Heading 4": (11, DARK_BLUE, 8, 4),
        "Caption": (9, MUTED, 4, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name.startswith("Heading") or name == "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = name.startswith("Heading")

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("信息解析 Skills 汇编"), size=9, color=MUTED)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_run_font(fp.add_run("第 "), size=9, color=MUTED)
    add_field(fp, "PAGE")
    set_run_font(fp.add_run(" 页"), size=9, color=MUTED)


def add_inline_markdown(paragraph, text: str, size=11, color=INK) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]), size=size, color=color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(9, size - 0.5), color=DARK_BLUE, latin="Menlo", east_asia="Microsoft YaHei")
            run.font.highlight_color = None
        elif token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), size=size, bold=True, color=color)
        else:
            label, target = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(label)
            set_run_font(run, size=size, color=BLUE)
            run.underline = True
            # Preserve the Markdown link target because python-docx does not expose hyperlinks directly.
            set_run_font(paragraph.add_run(f"（{target}）"), size=max(9, size - 1), color=MUTED)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=size, color=color)


def strip_frontmatter(text: str) -> tuple[dict[str, str], list[str]]:
    lines = text.splitlines()
    metadata: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        end = next((i for i in range(1, len(lines)) if lines[i].strip() == "---"), None)
        if end is not None:
            for line in lines[1:end]:
                if ":" in line and not line.startswith(" "):
                    key, value = line.split(":", 1)
                    metadata[key.strip()] = value.strip().strip('"')
            lines = lines[end + 1 :]
    while lines and not lines[0].strip():
        lines.pop(0)
    return metadata, lines


def is_table_separator(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int] | None:
    if start + 1 >= len(lines) or "|" not in lines[start] or not is_table_separator(lines[start + 1]):
        return None
    rows = []
    i = start
    while i < len(lines) and "|" in lines[i] and lines[i].strip():
        if i != start + 1:
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    return rows, i


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    widths = [9360 // cols] * cols
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_markdown(p, value, size=9.5, color=INK)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, code: list[str], language: str = "") -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    set_cell_margins(cell, top=100, bottom=100, start=140, end=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if language:
        run = p.add_run(f"[{language}]\n")
        set_run_font(run, size=8.5, bold=True, color=MUTED, latin="Menlo")
    run = p.add_run("\n".join(code))
    set_run_font(run, size=8.5, color=INK, latin="Menlo", east_asia="Microsoft YaHei")
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_markdown_body(doc: Document, lines: list[str]) -> None:
    # Skip the source document's first H1 because the compilation provides a numbered chapter title.
    skipped_first_h1 = False
    i = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline_markdown(p, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            language = stripped[3:].strip()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            add_code_block(doc, code, language)
            i += 1
            continue
        parsed = parse_table(lines, i)
        if parsed:
            flush_paragraph()
            rows, i = parsed
            add_markdown_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            if level == 1 and not skipped_first_h1:
                skipped_first_h1 = True
                i += 1
                continue
            style_level = min(4, max(2, level + 1))
            p = doc.add_paragraph(style=f"Heading {style_level}")
            add_inline_markdown(p, heading.group(2), size={2: 13, 3: 12, 4: 11}[style_level], color=BLUE if style_level < 4 else DARK_BLUE)
            i += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if bullet or numbered:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline_markdown(p, (bullet or numbered).group(1))
            i += 1
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_inline_markdown(p, stripped.lstrip("> "), color=MUTED)
            i += 1
            continue
        if stripped in ("---", "***", "___"):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            p_pr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:color"), "D0D5DD")
            borders.append(bottom)
            p_pr.append(borders)
            i += 1
            continue
        if not stripped:
            flush_paragraph()
        else:
            paragraph_buffer.append(line)
        i += 1
    flush_paragraph()


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    set_run_font(kicker.add_run("技术参考手册"), size=11, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("信息解析 Skills 汇编"), size=28, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("论文数据解析、PDF 元数据抽取、清洗与作者角色识别"), size=14, color=MUTED)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(18)
    rule.paragraph_format.space_after = Pt(26)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), BLUE)
    borders.append(bottom)
    p_pr.append(borders)

    names = doc.add_paragraph()
    names.alignment = WD_ALIGN_PARAGRAPH.CENTER
    names.paragraph_format.space_after = Pt(36)
    set_run_font(names.add_run("conference-paper-extraction\nconference-grobid-extract-metadata\nconference-clean-paper-metadata\nconference-author-role-extraction"), size=10.5, color=INK)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run(f"整理日期：{date.today().isoformat()}\n来源仓库：llm_conference_skill"), size=10, color=MUTED)
    doc.add_page_break()


def add_overview(doc: Document) -> None:
    doc.add_heading("文档说明", level=1)
    p = doc.add_paragraph()
    add_inline_markdown(p, "本文件将信息解析阶段的四个 Skill 统一整理为一份技术参考手册。各章节保留原 SKILL.md 的完整正文、命令示例、字段约束与工作流程；封面、总览和章节编号为汇编时新增。")

    doc.add_heading("处理链路", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    for idx, label in enumerate(("原始 JSON", "论文元数据", "PDF / TEI 证据", "清洗后事实", "作者角色事实")):
        if idx:
            set_run_font(p.add_run("  →  "), size=12, bold=True, color=BLUE)
        set_run_font(p.add_run(label), size=10.5, bold=True, color=INK)

    doc.add_heading("Skill 总览", level=2)
    table = doc.add_table(rows=1 + len(SOURCES), cols=3)
    table.style = "Table Grid"
    headers = ("Skill", "定位", "主要职责")
    for idx, text in enumerate(headers):
        set_cell_shading(table.cell(0, idx), LIGHT_BLUE)
        p = table.cell(0, idx).paragraphs[0]
        set_run_font(p.add_run(text), size=9.5, bold=True, color=INK)
    for r_idx, (folder, role, summary) in enumerate(SOURCES, 1):
        values = (folder, role, summary)
        for c_idx, value in enumerate(values):
            p = table.cell(r_idx, c_idx).paragraphs[0]
            set_run_font(p.add_run(value), size=9.5, color=INK)
    set_table_geometry(table, [2700, 1900, 4760])

    doc.add_heading("阅读建议", level=2)
    for text in (
        "先阅读每章开头的定位说明，再根据任务类型进入对应工作流。",
        "涉及 PDF 的字段冲突时，以 PDF 首页、TEI 和提取证据为主要复核依据。",
        "机构名称规范化不在本文件范围内，应在信息解析完成后交给 conference-affiliation-normalizer。",
    ):
        p = doc.add_paragraph(style="List Bullet")
        add_inline_markdown(p, text)


def add_source_chapter(doc: Document, index: int, folder: str, role: str, summary: str) -> None:
    doc.add_page_break()
    path = ROOT / folder / "SKILL.md"
    metadata, lines = strip_frontmatter(path.read_text(encoding="utf-8"))

    heading = doc.add_heading(f"{index}. {folder}", level=1)
    heading.paragraph_format.page_break_before = False
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run(f"{role}｜"), size=11, bold=True, color=BLUE)
    set_run_font(p.add_run(summary), size=11, color=INK)

    meta_rows = [
        ("Skill 名称", metadata.get("name", folder)),
        ("源文件", str(path)),
        ("原始描述", metadata.get("description", "—")),
    ]
    table = doc.add_table(rows=len(meta_rows), cols=2)
    table.style = "Table Grid"
    for r_idx, (label, value) in enumerate(meta_rows):
        set_cell_shading(table.cell(r_idx, 0), LIGHT_GRAY)
        p1 = table.cell(r_idx, 0).paragraphs[0]
        set_run_font(p1.add_run(label), size=9.5, bold=True, color=INK)
        p2 = table.cell(r_idx, 1).paragraphs[0]
        set_run_font(p2.add_run(value), size=9.2, color=INK)
    set_table_geometry(table, [1800, 7560])

    doc.add_heading("原始 Skill 正文", level=2)
    add_markdown_body(doc, lines)


def add_source_index(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("附：源文件索引", level=1)
    p = doc.add_paragraph()
    add_inline_markdown(p, "以下路径为本次汇编使用的原始文件。Word 文档不会修改这些源文件。")
    for folder, _, _ in SOURCES:
        p = doc.add_paragraph(style="List Bullet")
        add_inline_markdown(p, str(ROOT / folder / "SKILL.md"))


def build() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc)
    add_overview(doc)
    for idx, source in enumerate(SOURCES, 1):
        add_source_chapter(doc, idx, *source)
    add_source_index(doc)
    doc.core_properties.title = "信息解析 Skills 汇编"
    doc.core_properties.subject = "四个信息解析 Skill 的统一技术参考手册"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Skill, 信息解析, GROBID, 元数据, 作者角色"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
