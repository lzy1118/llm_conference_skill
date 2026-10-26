from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


DOCX = Path("/Users/liziyu/Desktop/论文/基于Skill的学术会议论文元数据处理研究v4.docx")
MAX_FIGURE_4_HEIGHT_EMU = 2_670_556


def set_on_off(parent, tag: str, enabled: bool) -> None:
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    element.set(qn("w:val"), "1" if enabled else "0")


def repeat_table_header(row) -> None:
    set_on_off(row._tr.get_or_add_trPr(), "w:tblHeader", True)


def prevent_row_split(row) -> None:
    set_on_off(row._tr.get_or_add_trPr(), "w:cantSplit", True)


def previous_drawing(caption: Paragraph) -> Paragraph:
    element = caption._p.getprevious()
    while element is not None:
        if element.tag == qn("w:p") and element.xpath(".//w:drawing"):
            return Paragraph(element, caption._parent)
        element = element.getprevious()
    raise RuntimeError(f"No drawing found before caption: {caption.text}")


def constrain_drawing_height(paragraph: Paragraph, max_height: int) -> None:
    extents = paragraph._p.xpath(".//wp:extent | .//a:xfrm/a:ext")
    if not extents:
        raise RuntimeError("Drawing has no extent")

    current_height = int(extents[0].get("cy"))
    if current_height > max_height:
        factor = max_height / current_height
        for extent in extents:
            extent.set("cx", str(round(int(extent.get("cx")) * factor)))
            extent.set("cy", str(round(int(extent.get("cy")) * factor)))

    fmt = paragraph.paragraph_format
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    fmt.keep_with_next = True


def main() -> None:
    doc = Document(str(DOCX))

    figure_4_caption = next(
        paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith("图 4 ")
    )
    constrain_drawing_height(previous_drawing(figure_4_caption), MAX_FIGURE_4_HEIGHT_EMU)

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style and paragraph.style.name.lower() == "caption":
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = text.startswith("表")

    for table in doc.tables:
        if table.rows:
            repeat_table_header(table.rows[0])
        for row in table.rows:
            prevent_row_split(row)

    temp_path = DOCX.with_name(f"{DOCX.stem}.layout-final{DOCX.suffix}")
    doc.save(str(temp_path))
    temp_path.replace(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
