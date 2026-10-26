from pathlib import Path
import os
import re
import tempfile
import zipfile

from lxml import etree
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "李梓瑜-实践教学总结报告-模板.docx"
OUTPUT = ROOT / "李梓瑜-燕山大学全日制专业学位研究生专业实践教学总结报告-完成版.docx"


def set_run_font(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def configure_paragraph(paragraph, *, indent=True, heading=False):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.first_line_indent = Pt(24 if indent and not heading else 0)


def add_text_paragraph(cell, text, *, indent=True, heading=False):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, indent=indent, heading=heading)
    run = paragraph.add_run(text)
    set_run_font(run, size=12, bold=heading)
    return paragraph


def add_label_paragraph(cell, label, text):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, indent=False)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=12, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=12, bold=False)
    return paragraph


def clear_cell_content(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)


def allow_row_to_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    for child in list(tr_pr):
        if child.tag in {qn("w:cantSplit"), qn("w:trHeight")}:
            tr_pr.remove(child)


def build_summary(cell):
    clear_cell_content(cell)

    add_label_paragraph(
        cell,
        "实践总结：",
        "本报告依据论文《基于Skill的学术会议论文元数据自动化处理方法研究》及其项目资料整理，重点总结研究问题、技术流程、应用结果和实践认识。现有材料的实践时间记载不一致：叙述材料为2025.6—2025.12，模板预填为2025年9月1日至12月31日（18周）；本文不作推断，请以学校实践记录核实并统一。模板预填的实践地点、指导教师、实践形式及报告日期也请按实际记录核验。论文中的研究样本年份不等同于本人实践起止日期。",
    )
    add_label_paragraph(cell, "实践地点：", "燕山大学信息馆（以实际记录核实）")
    add_label_paragraph(cell, "实践时间：", "2025.6—2025.12（具体日期以学校实践记录为准）")
    add_label_paragraph(cell, "指导教师：", "张玉清（以实际指导关系核实）")
    add_label_paragraph(
        cell,
        "实践目的：",
        "本实践围绕学术会议论文元数据自动化处理展开，针对数据来源分散、PDF署名版式复杂、作者与机构对应关系难以恢复、机构名称和国家信息不统一等问题，梳理并总结一条由论文清单采集、PDF信息解析、作者—机构关系识别、机构实体规范化到计量分析的连续流程。项目以Skill组织处理程序、领域规则和阶段产物，目标是为跨会议、跨年份的机构、国家及合作分析提供结构一致、来源可追溯的数据基础。",
    )

    add_text_paragraph(cell, "实践内容：", indent=False, heading=True)

    add_text_paragraph(cell, "一、实践项目背景与任务要求", indent=False, heading=True)
    add_text_paragraph(
        cell,
        "学术会议论文信息分布于会议官网、文献数据库、出版平台和论文全文，各来源的字段、页面结构及数据完整程度并不一致。同一论文可能出现题名格式差异、作者顺序或机构字段缺失、机构缩写和历史名称并存、跨国分支及联合实验室混合署名等情况。尤其在PDF首页中，上标编号、脚注、多列布局和共享机构等版式会影响作者—机构关系识别。若不先处理记录对齐、署名关系和机构实体，就直接计算国家分布、机构参与量或合作网络，可能造成漏计、重复计数或实体关系错位。",
    )
    add_text_paragraph(
        cell,
        "本项目论文题为《基于Skill的学术会议论文元数据自动化处理方法研究》。研究范围为2020—2025年的七个安全与密码学会议：ASIACRYPT、CRYPTO、EUROCRYPT、NDSS、CCS、IEEE S&P和USENIX Security。需要完成的并非单一来源的论文检索，而是把不同来源的论文记录和PDF组织成可对齐输入，恢复论文、作者、机构及作者角色等信息，规范机构实体和国家字段，并依据清楚的计量口径生成统计结果。",
    )
    add_text_paragraph(
        cell,
        "报告根据论文与项目资料概括实践内容。论文描述的是项目方法和汇总结果，并未逐项说明研究生个人的任务分工，因此本文不将全部项目产出表述为个人独立实现；个人实际承担的具体工作可结合导师安排和过程记录补充确认。",
    )

    add_text_paragraph(cell, "二、基于Skill的处理框架与模块设计", indent=False, heading=True)
    add_text_paragraph(
        cell,
        "新版论文将流程按处理阶段组织为四类Skill：数据采集、信息解析、机构规范化和计量分析。每个Skill以任务说明、处理程序、规则资源和阶段输出构成执行单元，并通过统一论文标识及结构化中间结果衔接。数据采集阶段交付论文清单和整理后的PDF数据集；信息解析阶段形成论文、作者、原始机构及作者角色事实；机构规范化阶段生成可比较的机构与国家实体；计量分析阶段基于稳定事实表生成国家、机构、作者及合作关系结果。",
    )
    add_text_paragraph(
        cell,
        "其中，信息解析不是单一程序，而由基础论文记录构建、GROBID/TEI信息提取、论文记录清洗以及作者角色抽取等子任务衔接完成。Skill的作用是把会议入口、版式规则、机构别名和统计口径等任务知识与具体程序一并组织，明确各阶段输入、处理边界和输出契约。论文附录提供SKILL.md结构示例，用于说明模块组织方式；仅凭示例本身不推断所有脚本均已部署或投入生产。",
    )

    add_text_paragraph(cell, "三、多来源数据采集与PDF数据集构建", indent=False, heading=True)
    add_text_paragraph(
        cell,
        "数据采集以会议名称和年份配置入口，来源包括IACR结构化记录、DBLP、会议官网、出版平台和论文PDF。论文将七个会议在2020—2025年的研究范围按会议与年份组织为42个会议—年份单元，并从论文集目录、官网及出版页面整理题名、作者、年份、页码、DOI和页面地址等信息，形成目标论文清单。对于不同平台的入口差异，采集流程通过对应的来源解析规则处理，再向后续解析阶段提供统一索引。",
    )
    add_text_paragraph(
        cell,
        "PDF获取环节根据DOI、出版页面地址或平台标识访问论文详情页，并使用Selenium读取页面内容和PDF链接；需要页面会话状态的来源也在相应流程中处理。下载后检查文件类型、文件头和文件大小，以识别错误页面、空文件或不完整文件；同时计算SHA-256值识别内容重复项。最后统一题名中的大小写、标点和空格，对照目标论文清单完成PDF匹配。重复、未匹配或异常文件进入结果清单，不作为静默丢弃的记录。",
    )
    add_text_paragraph(
        cell,
        "该阶段形成按会议和年份组织的PDF数据集，以及包含题名、来源页面、PDF地址、本地路径、匹配状态和文件哈希等信息的采集清单。保留来源和校验信息，使后续解析结果能够回到对应论文记录和文件核对，也便于区分论文清单、有效PDF及后续关系事实等不同数据层次。",
    )

    add_text_paragraph(cell, "四、PDF信息解析与作者—机构关系恢复", indent=False, heading=True)
    add_text_paragraph(
        cell,
        "元数据提取以GROBID生成的header TEI作为结构化起点，结合PDF前两页的文本内容和版面信息，提取并核对题名、作者、机构、邮箱、ORCID及作者角色。论文方法利用文本块的纵向位置、横向位置和字号识别题名、作者与机构区域；当GROBID字段缺失、截断或混入页面噪声时，再使用版面候选补充或校正。作者顺序以融合后的结果为准，字段融合保留来源证据，避免仅凭连续文本顺序推断作者与机构关系。",
    )
    add_text_paragraph(
        cell,
        "针对复杂署名，论文整理了七类版式规则：上标编号、作者信息块、横向矩阵、纵向堆叠、分组共享、全体共享以及脚注与跨页补充。规则分别依据编号或符号对应、局部信息块、列对齐、上下文边界、共享范围和脚注位置建立关系候选。一名作者可对应多个机构，多名作者也可共享同一机构；无法确定唯一归属时不强行合并。第一作者依据最终作者顺序识别，共同一作和通讯作者则要求有明确姓名、符号或角色说明。",
    )
    add_text_paragraph(
        cell,
        "当PDF文本层丢失上标、符号或机构编号时，论文还以会议官方页面DOM作为关系补充来源。以CCS 2025论文“​​Nebula: Efficient, Private and Accurate Histogram Estimation​​”为例，页面节点按“作者（机构）”组织署名；解析过程在单篇论文节点内识别作者和机构单元，按原署名顺序保存关系，并把复合署名交由后续规范化阶段处理。网页证据与PDF中提取的作者顺序、邮箱、ORCID及角色信息按论文标识和规范化姓名进行对齐。",
    )

    add_text_paragraph(cell, "五、机构名称规范化与统计口径", indent=False, heading=True)
    add_text_paragraph(
        cell,
        "机构署名可能同时包含主体机构、院系、实验室、地址和国家，也可能以全称、简称、历史名称或不同层级表达。论文采用AffilGood解析原始affiliation，识别机构主体、下属单位、城市和国家，再利用ROR中的规范名称、别名、缩写及地理信息进行实体链接。未能唯一匹配的记录进入候选检索：从ROR召回候选后结合国家、城市信息和多语言重排结果筛选；论文还描述了在不确定候选范围内使用Qwen2.5-3B-Instruct辅助判别的配置。低于既定阈值或被拒绝的记录保留原始署名和未匹配状态。经审核的本地机构词典用于补充缺失字段和确认的名称变体，不覆盖已接受的ROR链接。",
    )
    add_text_paragraph(
        cell,
        "规范化结果同时保留原始机构署名、规范名称、机构标识、国家及匹配方式。实体归并以ROR或项目内部稳定标识为依据；历史名称仅在组织延续关系明确时合并，跨国分支和具有独立组织属性的单位分别保留。该处理有助于在统一统计实体的同时维持原始署名证据。",
    )
    add_text_paragraph(
        cell,
        "计量分析以不同事实关系分别聚合：论文数按去重后的独立论文计算；国家和机构参与量按论文—实体关系去重，同一论文中同一实体只计一次，但一篇论文可以涉及多个实体，因此各国或机构参与量及占比不能视作互斥份额。作者统计按作者—论文关系去重；合作网络按单篇论文中实体共现生成无向边，并按共同署名论文数累计边权。稳定实体标识用于统计，展示名称用于结果呈现。",
    )

    add_text_paragraph(cell, "六、项目结果与完成情况", indent=False, heading=True)
    add_text_paragraph(
        cell,
        "论文报告的应用数据覆盖七个会议、2020—2025年共7,079篇论文，并形成35,680条作者—机构关系记录。年度论文数依次为2020年721篇、2021年920篇、2022年986篇、2023年1,289篇、2024年1,521篇和2025年1,642篇。35,680是作者—机构关系记录数，不是唯一作者人数；7,079是研究范围内的论文数，两者属于不同统计单位。",
    )
    add_text_paragraph(
        cell,
        "机构参与统计中，CISPA Helmholtz Center for Information Security参与351篇论文，Tsinghua University参与335篇，ETH Zurich参与248篇；Zhejiang University和Purdue University分别参与233篇和228篇。国家参与统计中，美国、中国和德国分别参与3,763篇、1,613篇和1,013篇论文，占全部论文的53.16%、22.79%和14.31%；瑞士和法国分别为509篇和471篇。由于一篇论文可以计入多个国家或机构，以上比例为参与度统计，不应相加解释为互斥的国家份额。",
    )
    add_text_paragraph(
        cell,
        "机构合作网络包含2,247个机构节点和13,326条合作边，平均度为11.86。其中10,074条边对应1篇共同署名论文，3,252条边对应2篇及以上论文，至少共同署名10篇论文的机构组合有145组。Tsinghua University与Zhongguancun Laboratory共同署名108篇论文；国家组合中，中国与美国共同参与603篇。此处13,326是机构合作网络边数，不是作者节点或作者合作边。",
    )
    add_text_paragraph(
        cell,
        "这些结果展示了标准化数据用于国家、机构和合作关系分析的方式，也体现出机构标识和计量口径对结果解释的重要性。论文提供了流程设计、数据规模和应用统计，但未报告抽取准确率、召回率、F1、基线对比或运行性能等定量质量指标。因此，论文中的样本规模和统计结果不能替代对解析正确性的独立评估，也不应据此宣称方法达到某一准确率或相对基线显著提升。",
    )

    add_text_paragraph(cell, "七、实践中需要处理的问题与方法", indent=False, heading=True)
    add_text_paragraph(
        cell,
        "①来源入口与记录完整度存在差异。会议、年份和出版平台的页面结构并不相同，某些字段还需跨来源对齐。处理时按会议和年份管理入口，保留题名、页面、链接、文件哈希和匹配状态，并将重复、未匹配或异常文件写入清单，使采集结果可核查。",
    )
    add_text_paragraph(
        cell,
        "②PDF文本抽取不等同于署名关系恢复。复杂版式中的上标、分栏、共享机构和脚注容易在纯文本中丢失或错序。论文采用GROBID结构结果与页面坐标相结合的策略，并为七类常见署名版式配置规则；对于不能唯一判定的候选保留不确定状态，必要时利用官方页面DOM补充关系，而不是把文本邻近直接当作作者—机构对应。",
    )
    add_text_paragraph(
        cell,
        "③机构实体匹配需要在归并与误合并之间权衡。缩写、别名和历史名称可能指向同一机构，但院系、分支或联合实验室也可能具有独立组织属性。AffilGood与ROR提供候选实体和规范标识，本地词典补充经过核实的映射；信息不足时保留原始署名和未匹配状态。机构参与数及网络结构受实体链接结果影响，解读时需要同时考虑匹配证据和仍未确认的记录。",
    )
    add_text_paragraph(
        cell,
        "④项目结果边界和统计分母需要清楚说明。论文规模、作者—机构关系行数、国家参与论文数和网络边数分别对应不同对象。尤其是国家和机构参与统计允许一篇论文关联多个实体，不能将它们解释为互斥份额；网络边表示机构之间的共现关系，也不等于共同论文篇数本身。通过在统计前确定去重单位、实体标识和边权语义，可以减少同名字段或重复关系对结论的影响。",
    )
    add_text_paragraph(
        cell,
        "⑤自动化流程仍需要可复核和持续维护。出版页面、机构注册信息和名称变体可能变化，Skill说明和结构示例也不能代替对实际运行结果的验证。保留原始记录、阶段输出、匹配状态和规则依据，有助于后续定位问题；对于准确率、吞吐、成本和跨学科泛化等尚未量化的方面，应作为后续验证工作，而不作超出论文证据的结论。",
    )

    add_text_paragraph(cell, "八、实践心得体会", indent=False, heading=True)
    add_text_paragraph(
        cell,
        "通过本次专业实践，我对科研计量分析所依赖的数据基础有了更具体的认识。分析结果不只取决于后续统计方法，还受到论文清单、来源证据、PDF版面解析、作者—机构对应以及实体规范化等前置环节影响。新版论文将这些环节连接为数据采集、信息解析、机构规范化和计量分析四类Skill，使我认识到复杂研究任务需要明确阶段边界和输出契约，才能便于复核与持续维护。",
    )
    add_text_paragraph(
        cell,
        "本次实践也加深了我对非结构化信息解析和实体链接的理解。GROBID提供的字段结果需要结合PDF版面证据校核；网页DOM能够补充部分关系，但仍须保存其来源和解析规则；AffilGood与ROR有助于机构匹配，却不能消除所有模糊情况。保留原始署名、规范标识和未匹配状态，比对所有输入强行给出确定结果更有利于解释数据边界。",
    )
    add_text_paragraph(
        cell,
        "此外，我进一步理解了统计口径在结果表达中的作用。7,079篇论文与35,680条作者—机构关系、国家参与量与合作网络边数分别是不同的度量单位；如果不说明计数对象，读者容易把关系记录误认为人数、把参与率误认为互斥份额。今后整理和汇报研究结果时，我会先核对数据对象、分母和去重规则，再解释统计现象，并将方法实现证据与应用结果区分开来。",
    )
    add_text_paragraph(
        cell,
        "论文目前的应用范围集中于七个安全与密码学会议，后续仍需扩展会议和出版平台规则，并继续完善历史名称、分支机构及联合实验室的处理。对于解析质量和运行性能，也需要通过人工标注样本、基线比较及运行测试进行独立评估。总体而言，本次实践帮助我从单纯关注最终统计结果，转向同时关注数据来源、关系证据、处理边界和结论可复核性，为后续科研与工程实践积累了方法认识。",
    )

    add_text_paragraph(cell, "研究生本人签字：             ", indent=False, heading=True)
    add_text_paragraph(cell, "                                            ", indent=False, heading=True)
    add_text_paragraph(cell, "校内导师签字：             ", indent=False, heading=True)
    add_text_paragraph(cell, "校外导师签字：             ", indent=False, heading=True)


def widen_footer_page_number_frame(docx_path):
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".report-", suffix=".docx", dir=docx_path.parent, delete=False
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        namespaces = {
            "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            "v": "urn:schemas-microsoft-com:vml",
        }
        with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(
            temporary_path, "w"
        ) as destination:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/footer1.xml":
                    root = etree.fromstring(data)
                    for extent in root.xpath(".//wp:extent | .//a:ext", namespaces=namespaces):
                        if int(extent.get("cx", "0")) < 254000:
                            extent.set("cx", "254000")
                    for rect in root.xpath(".//v:rect", namespaces=namespaces):
                        style = rect.get("style", "")
                        style = re.sub(r"(?<![A-Za-z-])width:\s*[^;]+pt", "width:20pt", style, count=1)
                        rect.set("style", style)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(item, data)
        os.replace(temporary_path, docx_path)
        temporary_path = None
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def main():
    document = Document(TEMPLATE)
    table = document.tables[0]
    build_summary(table.cell(2, 0))
    allow_row_to_split(table.rows[2])
    document.save(OUTPUT)
    widen_footer_page_number_frame(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
